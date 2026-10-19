"""Regenerates the PDF, DOCX and TXT test fixtures in this folder.

    python make_documents.py

Needs reportlab, pypdf, python-docx and Pillow, and LibreOffice (`soffice`) on the PATH
for the Arabic PDF, whose text needs real shaping. Every file stays small.
"""

import io
import pathlib
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from PIL import Image
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

HERE = pathlib.Path(__file__).resolve().parent


def quadrant_png():
    """The same four-colour picture as the image fixtures, larger."""
    img = Image.new("RGB", (160, 120))
    px = img.load()
    colours = [(220, 30, 30), (30, 200, 60), (40, 60, 220), (250, 250, 250)]
    for y in range(120):
        for x in range(160):
            px[x, y] = colours[(1 if x >= 80 else 0) + (2 if y >= 60 else 0)]
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def text_pdf():
    """Two A4 pages of text, a picture on the second, and document properties."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, invariant=1)
    c.setTitle("Quarterly report")
    c.setAuthor("Test Author")
    c.setSubject("Dropwheel fixture")
    c.setKeywords("alpha, beta")
    c.setCreator("make_documents.py")
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 760, "Quarterly report")
    c.setFont("Helvetica", 12)
    lines = [
        "Dropwheel turns this page into text.",
        "Second line of the first paragraph.",
        "",
        "A new paragraph after a blank line.",
    ]
    y = 720
    for line in lines:
        c.drawString(72, y, line)
        y -= 16
    c.showPage()
    c.setFont("Helvetica", 12)
    c.drawString(72, 760, "Page two has a picture below.")
    c.drawImage(ImageReader(io.BytesIO(quadrant_png())), 72, 500, width=320, height=240)
    c.showPage()
    c.save()
    (HERE / "text.pdf").write_bytes(buf.getvalue())


def scan_pdf():
    """One page that is only a picture, like a scan."""
    img = Image.open(io.BytesIO(quadrant_png()))
    buf = io.BytesIO()
    img.save(buf, "PDF", resolution=36.0)
    (HERE / "scan.pdf").write_bytes(buf.getvalue())


def rotated_pdf():
    """A portrait page shown in landscape through /Rotate 90, with a mark in a corner."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(200, 300), invariant=1)
    c.setFillColorRGB(1, 0, 0)
    # Bottom left of the unrotated page; top left once turned clockwise.
    c.rect(0, 0, 50, 50, stroke=0, fill=1)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 10)
    c.drawString(60, 150, "Rotated")
    c.showPage()
    c.save()
    reader = PdfReader(io.BytesIO(buf.getvalue()))
    writer = PdfWriter()
    page = reader.pages[0]
    page.rotate(90)
    writer.add_page(page)
    out = io.BytesIO()
    writer.write(out)
    (HERE / "rotated.pdf").write_bytes(out.getvalue())


def protected_pdf():
    reader = PdfReader(HERE / "text.pdf")
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password="secret", owner_password="owner", algorithm="AES-128")
    out = io.BytesIO()
    writer.write(out)
    (HERE / "protected.pdf").write_bytes(out.getvalue())


def broken_pdf():
    data = (HERE / "text.pdf").read_bytes()
    (HERE / "broken.pdf").write_bytes(data[: len(data) // 3])


def set_rtl(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = ppr.makeelement(qn("w:bidi"), {})
    ppr.append(bidi)
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rpr.append(rpr.makeelement(qn("w:rtl"), {}))


def report_docx():
    """Headings, styled runs, lists, a table, a picture, a page break and Arabic."""
    doc = Document()
    props = doc.core_properties
    props.title = "Quarterly report"
    props.author = "Test Author"
    props.subject = "Dropwheel fixture"
    props.keywords = "alpha, beta"
    props.comments = "Made by make_documents.py"
    props.last_modified_by = "Someone Else"

    doc.add_heading("Quarterly report", level=1)
    p = doc.add_paragraph("Plain text, then ")
    p.add_run("bold").bold = True
    p.add_run(", ")
    p.add_run("italic").italic = True
    p.add_run(" and ")
    r = p.add_run("red")
    r.font.color.rgb = RGBColor(0xC0, 0x10, 0x10)
    p.add_run(" words.")
    doc.add_heading("Lists", level=2)
    doc.add_paragraph("First point", style="List Bullet")
    doc.add_paragraph("Second point", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    doc.add_paragraph("Step two", style="List Number")
    centred = doc.add_paragraph("Centred line")
    centred.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for i, text in enumerate(["Region", "Sales", "Change"]):
        table.cell(0, i).text = text
    for i, text in enumerate(["North", "1,200", "+4%"]):
        table.cell(1, i).text = text

    doc.add_paragraph("A picture:")
    doc.add_picture(io.BytesIO(quadrant_png()), width=Cm(4))

    brk = doc.add_paragraph()
    brk.add_run().add_break(WD_BREAK.PAGE)
    arabic = doc.add_paragraph("مرحبا بالعالم. هذا تقرير ربع سنوي.")
    set_rtl(arabic)
    doc.add_paragraph("Last line.")
    doc.save(HERE / "report.docx")


def arabic_pdf():
    """An Arabic page laid out by LibreOffice, which shapes Arabic properly."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("skipping arabic.pdf: LibreOffice not found")
        return
    with tempfile.TemporaryDirectory() as tmp:
        tmp = pathlib.Path(tmp)
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(14)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        fonts.set(qn("w:cs"), "DejaVu Sans")
        p = doc.add_paragraph("مرحبا بالعالم")
        set_rtl(p)
        doc.add_paragraph("Hello world")
        src = tmp / "arabic.docx"
        doc.save(src)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
            check=True,
            capture_output=True,
            env={"HOME": str(tmp), "PATH": "/usr/bin:/bin"},
        )
        # Drop LibreOffice's producer and dates so the fixture is reproducible-ish.
        reader = PdfReader(tmp / "arabic.pdf")
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        out = io.BytesIO()
        writer.write(out)
        (HERE / "arabic.pdf").write_bytes(out.getvalue())


def texts():
    (HERE / "notes.txt").write_text(
        "Shopping list\n\tmilk\n\tbread\n\nملاحظة: لا تنسَ الخبز\nLast line without newline",
        encoding="utf-8",
    )
    (HERE / "legacy.txt").write_bytes("مرحبا بالعالم\r\nهذا ملف قديم\r\n".encode("cp1256"))
    (HERE / "utf16.txt").write_bytes("Ünïcode in UTF-16\r\nسطر ثان\r\n".encode("utf-16"))


if __name__ == "__main__":
    text_pdf()
    scan_pdf()
    rotated_pdf()
    protected_pdf()
    broken_pdf()
    report_docx()
    arabic_pdf()
    texts()
    for f in sorted(HERE.iterdir()):
        if f.suffix in {".pdf", ".docx", ".txt"}:
            print(f"{f.name}: {f.stat().st_size} bytes")
